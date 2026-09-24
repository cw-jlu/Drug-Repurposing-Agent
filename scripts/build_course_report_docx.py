"""Build an editable Word draft from the verified Chinese report manuscript."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("docs/课程设计报告_草稿.md")
OUTPUT = Path("deliverables/药物重定位Agent_课程设计报告草稿.docx")
FONT = "Microsoft YaHei"


def set_east_asian_font(style) -> None:
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), FONT)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url,
                         "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174F78")
    props.append(color)
    run.append(props)
    txt = OxmlElement("w:t")
    txt.text = label
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)


TOKEN = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)")


def add_inline(paragraph, content: str) -> None:
    parts = TOKEN.split(content)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = part[1:-1].split("](", 1)
            add_hyperlink(paragraph, label, url)
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1])
        else:
            paragraph.add_run(part)


def heading_text(content: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", content)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_borders(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)


def set_cell_margin(cell, twips: int = 75) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:w"), str(twips))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    tcpr.append(margins)


def add_table(document: Document, rows: list[list[str]]) -> None:
    width = len(rows[0])
    table = document.add_table(rows=len(rows), cols=width)
    table.autofit = False
    if width == 3:
        widths = [Inches(2.45), Inches(2.2), Inches(2.2)]
    else:
        widths = [Inches(6.85 / width)] * width
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Normal"]
            paragraph.paragraph_format.space_after = Pt(0)
            if j > 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, value.strip())
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if i == 0:
                set_cell_shading(cell, "18384E")
            elif i % 2 == 0:
                set_cell_shading(cell, "F2F6F8")
            set_cell_borders(cell)
            set_cell_margin(cell)
    header_trpr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_trpr.append(repeat)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.34)
    section.footer_distance = Inches(0.34)

    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        set_east_asian_font(styles[name])
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    normal = styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(5)
    styles["Title"].font.size = Pt(18)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(10)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "药物重定位 Agent 课程设计报告草稿"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = FONT
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 100, 108)
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", value) for value in values):
                    rows.append(values)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline(p, line[2:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, heading_text(line[3:]))
            if "数据来源与质量控制" in line:
                p.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, heading_text(line[4:]))
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.12)
            add_inline(p, line[2:])
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(90, 100, 108)
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
